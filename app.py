import streamlit as st
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google import genai
from google.genai import types
import datetime
import io
import json
import os
from pathlib import Path

st.set_page_config(page_title="Automazione Notifiche e Atti", layout="wide")

BASE_DIR = Path(__file__).resolve().parent
MODELLI_DIR = BASE_DIR / "modelli_riferimento"
MODELLI_DIR.mkdir(parents=True, exist_ok=True)
RUBRICA_FILE = BASE_DIR / "enti_rubrica.json"

if "reset_count" not in st.session_state:
    st.session_state["reset_count"] = 0
if "dati_elaborati" not in st.session_state:
    st.session_state["dati_elaborati"] = None

def get_api_key():
    if "GEMINI_API_KEY" in st.secrets:
        return st.secrets["GEMINI_API_KEY"]
    if "gemini_key" in st.session_state:
        return st.session_state["gemini_key"]
    return ""

api_key = get_api_key()

with st.sidebar:
    st.title("Impostazioni")
    if not api_key:
        api_input = st.text_input("Gemini API Key:", type="password")
        if api_input:
            st.session_state["gemini_key"] = api_input
            api_key = api_input
            st.rerun()
    else:
        st.success("API Key configurata!")
    
    st.subheader("Modelli Word")
    nuovo_modello = st.file_uploader("Carica modello (.docx)", type=["docx"], key="side_modello")
    if nuovo_modello:
        with open(MODELLI_DIR / nuovo_modello.name, "wb") as f:
            f.write(nuovo_modello.getbuffer())
        st.success("Modello salvato!")
        st.rerun()

def carica_rubrica():
    if RUBRICA_FILE.exists():
        try:
            with open(RUBRICA_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def aggiorna_rubrica(nome_ente, blocco_indirizzo):
    if not nome_ente or not blocco_indirizzo:
        return
    rubrica = carica_rubrica()
    rubrica[nome_ente.strip()] = blocco_indirizzo.strip()
    try:
        with open(RUBRICA_FILE, "w", encoding="utf-8") as f:
            json.dump(rubrica, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

st.title("Generatore Lettere di Trasmissione")

oggi = datetime.date.today()
data_odierna_str = oggi.strftime("%d/%m/%Y")
anno_corrente = oggi.year

modelli_docx = [f for f in os.listdir(MODELLI_DIR) if f.endswith(".docx") and not f.startswith("~$")]
opzioni_modello = ["-- Modello Istituzionale Standard (Generato al volo) --"] + modelli_docx
modello_scelto = st.selectbox("Formato / Modello Word di base:", opzioni_modello)

col_prot, col_data = st.columns(2)
with col_prot:
    protocollo_input = st.text_input("Protocollo Uscita (in alto a sinistra)", value=f"N. /{anno_corrente}")
with col_data:
    st.info(f"📅 **Data documento (in alto a destra):** {data_odierna_str}")

rubrica_enti = carica_rubrica()
opzioni_enti = ["-- Rileva automaticamente dall'atto caricato --"] + list(rubrica_enti.keys())
ente_selezionato = st.selectbox("Destinatario da Rubrica:", opzioni_enti)

st.subheader("1. Atto/Richiesta di Notifica ricevuta")

col_btn_reset, _ = st.columns([1, 4])
with col_btn_reset:
    if st.button("🔄 Pulisci/Sblocca Uploader"):
        st.session_state["reset_count"] += 1
        st.session_state["dati_elaborati"] = None
        st.rerun()

modalita_caricamento = st.radio(
    "Scegli come inserire l'atto:",
    ["Carica File (PDF, Foto, Word)", "Scatta Foto con Fotocamera (da cellulare/webcam)"],
    horizontal=True
)

file_atto = None
if modalita_caricamento == "Carica File (PDF, Foto, Word)":
    file_atto = st.file_uploader(
        "Seleziona o trascina il file dell'atto:",
        key=f"file_uploader_{st.session_state['reset_count']}"
    )
else:
    file_atto = st.camera_input("Inquadra e scatta la foto all'atto:")

if file_atto:
    st.success(f"📎 Documento agganciato: **{getattr(file_atto, 'name', 'Foto scattata')}** ({round(len(file_atto.getvalue())/1024, 1)} KB)")
else:
    st.info("ℹ️ Nessun documento agganciato. Carica un file o scatta una foto prima di procedere.")

st.subheader("2. Note operative")
note_input = st.text_input(
    "Dettagli operativi (opzionale):",
    placeholder="Es. Notificato a mani proprie / irreperibile..."
)

def genera_docx_standard(data_str, prot_uscita, prot_rif, destinatari, oggetto, corpo):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    p_top = doc.add_paragraph()
    r_prot = p_top.add_run(f"Prot. {prot_uscita}")
    r_prot.bold = True
    r_prot.font.size = Pt(10)
    p_top.add_run("\t\t\t\t\t\t")
    r_data = p_top.add_run(f"Data: {data_str}")
    r_data.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_before = Pt(18)

    p_dest = doc.add_paragraph()
    p_dest.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for riga in destinatari.split("\n"):
        if riga.strip():
            r = p_dest.add_run(riga.strip() + "\n")
            r.bold = True
            r.font.size = Pt(11)

    doc.add_paragraph().paragraph_format.space_before = Pt(14)

    p_ogg = doc.add_paragraph()
    r_ogg_l = p_ogg.add_run("OGGETTO: ")
    r_ogg_l.bold = True
    r_ogg = p_ogg.add_run(oggetto)
    r_ogg.bold = True

    doc.add_paragraph().paragraph_format.space_before = Pt(14)

    p_corpo = doc.add_paragraph()
    p_corpo.paragraph_format.line_spacing = 1.15
    p_corpo.paragraph_format.space_after = Pt(12)
    p_corpo.add_run(corpo)

    doc.add_paragraph().paragraph_format.space_before = Pt(36)

    p_firma = doc.add_paragraph()
    p_firma.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_firma.add_run("IL COMANDANTE\n(firma e timbro)")

    if prot_rif:
        doc.add_paragraph().paragraph_format.space_before = Pt(30)
        p_rif = doc.add_paragraph()
        r_rif = p_rif.add_run(prot_rif)
        r_rif.font.size = Pt(9)
        r_rif.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def elabora_con_gemini(chiave, file_caricato, note_op, indirizzo_fisso=None):
    client = genai.Client(api_key=chiave)
    contenuti = []

    prompt = f"""
Sei un addetto esperto delle Forze dell'Ordine / Pubblica Amministrazione italiana incaricato di redigere la lettera di trasmissione per la restituzione di un atto notificato.

REGOLE TASSATIVE:
1. PROTOCOLLO IN BASSO A SINISTRA:
   - Estrai protocollo, R.G.N.R., R.G. DIB o estremi dell'atto arrivato (es. "Rif. nota prot. n. 1234 del 12/03/2026").

2. DESTINATARI:
   - Individua l'autorità o ufficio mittente (Tribunale, Procura, Questura, Prefettura, ecc.) con cancelleria e indirizzo/PEC.
   - Fornisci un "nome_breve_ente" sintetico per la rubrica.

3. DATI ANAGRAFICI ED OGGETTO:
   - Estrai: cognome, nome, data e luogo di nascita, residenza e domicilio (se presenti).
   - L'oggetto DEVE essere formulato nel seguente modo:
     "Trasmissione atti notificati a carico di [COGNOME Nome, nato a LUOGO (PROV) il GG/MM/AAAA, residente a COMUNE (PROV) in VIA/PIAZZA, domiciliato a COMUNE (PROV) in VIA/PIAZZA]"

4. CORPO DELLA LETTERA:
   - DEVE iniziare tassativamente con:
     "Allegato alla presente si restituisce debitamente notificato al soggetto sopra meglio indicato "
   - Prosegui con il tipo di atto e i riferimenti precisi, terminando con "per il prosieguo di competenza.".

NOTE AGGIUNTIVE: {note_op if note_op else "Notifica eseguita."}

Rispondi ESCLUSIVAMENTE con un oggetto JSON valido (senza testo prima o dopo) con queste chiavi:
{{
  "protocollo_riferimento": "Rif. prot. ...",
  "nome_breve_ente": "Nome per rubrica",
  "destinatari": "Ente e Ufficio\\nIndirizzo/PEC",
  "oggetto": "Oggetto completo",
  "corpo_lettera": "Allegato alla presente si restituisce debitamente notificato..."
}}
"""
    contenuti.append(prompt)

    b = file_caricato.getvalue()
    nome = getattr(file_caricato, "name", "atto.jpg").lower()
    mime = getattr(file_caricato, "type", "") or ""

    if nome.endswith(".pdf") or "pdf" in mime:
        contenuti.append(types.Part.from_bytes(data=b, mime_type="application/pdf"))
    elif any(nome.endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp"]) or "image" in mime:
        tipo_img = mime if "image" in mime else "image/jpeg"
        contenuti.append(types.Part.from_bytes(data=b, mime_type=tipo_img))
    elif nome.endswith(".docx"):
        tdoc = Document(io.BytesIO(b))
        txt = "\n".join([p.text for p in tdoc.paragraphs if p.text.strip()])
        contenuti.append(f"\nTESTO ATTO:\n{txt}")
    else:
        contenuti.append(types.Part.from_bytes(data=b, mime_type="application/pdf"))

    # Configurazione corretta per google-genai
    config = types.GenerateContentConfig(
        response_mime_type="application/json"
    )

    # Chiamata al modello con fallback
    modelli = ["gemini-2.5-flash", "gemini-2.0-flash"]
    res = None
    ultimo_errore = None

    for m in modelli:
        try:
            res = client.models.generate_content(
                model=m,
                contents=contenuti,
                config=config
            )
            if res and res.text:
                break
        except Exception as e:
            ultimo_errore = e
            continue

    if not res or not res.text:
        raise Exception(f"Errore chiamata Gemini: {ultimo_errore}")

    t = res.text.strip()
    if t.startswith("```json"):
        t = t[7:]
    if t.startswith("```"):
        t = t[3:]
    if t.endswith("```"):
        t = t[:-3]
        
    dati = json.loads(t.strip())
    if indirizzo_fisso:
        dati["destinatari"] = indirizzo_fisso
    return dati

st.divider()

if st.button("🚀 Elabora e Genera Lettera di Trasmissione", type="primary", use_container_width=True):
    if not api_key:
        st.error("❌ Manca la Gemini API Key. Inseriscila nella barra laterale a sinistra o nei Secrets di Streamlit.")
    elif not file_atto:
        st.error("❌ Nessun documento caricato! Seleziona prima un file PDF/Foto o scatta una foto al punto 1.")
    else:
        with st.spinner("⏳ Analisi ed estrazione dati con Gemini in corso..."):
            try:
                ind_fisso = rubrica_enti.get(ente_selezionato) if not ente_selezionato.startswith("--") else None
                risultato = elabora_con_gemini(api_key, file_atto, note_input, ind_fisso)
                st.session_state["dati_elaborati"] = risultato

                if ente_selezionato.startswith("--") and risultato.get("nome_breve_ente") and risultato.get("destinatari"):
                    aggiorna_rubrica(risultato["nome_breve_ente"], risultato["destinatari"])

                st.success("✅ Atto analizzato con successo!")
            except Exception as e:
                st.error("Si è verificato un errore durante l'elaborazione:")
                st.exception(e)

if st.session_state.get("dati_elaborati"):
    dati = st.session_state["dati_elaborati"]
    st.subheader("📋 Dati Generati (Modificabili)")

    cp1, cp2 = st.columns(2)
    with cp1:
        prot_u = st.text_input("Protocollo Uscita:", value=protocollo_input)
    with cp2:
        prot_r = st.text_input("Protocollo/Rif. Atto:", value=dati.get("protocollo_riferimento", ""))

    c1, c2 = st.columns(2)
    with c1:
        dest_f = st.text_area("Destinatari / Ente:", value=dati.get("destinatari", ""), height=130)
    with c2:
        ogg_f = st.text_area("Oggetto Generato:", value=dati.get("oggetto", ""), height=130)

    corpo_f = st.text_area("Testo Trasmissione:", value=dati.get("corpo_lettera", ""), height=180)

    if modello_scelto.startswith("--"):
        word_buf = genera_docx_standard(data_odierna_str, prot_u, prot_r, dest_f, ogg_f, corpo_f)
    else:
        doc = Document(MODELLI_DIR / modello_scelto)
        mappa = {
            "{{DATA}}": data_odierna_str,
            "{{PROTOCOLLO}}": prot_u,
            "{{PROTOCOLLO_USCITA}}": prot_u,
            "{{PROTOCOLLO_RIFERIMENTO}}": prot_r,
            "{{RIFERIMENTO}}": prot_r,
            "{{DESTINATARI}}": dest_f,
            "{{OGGETTO}}": ogg_f,
            "{{CORPO_LETTERA}}": corpo_f
        }
        for p in doc.paragraphs:
            for k, v in mappa.items():
                if k in p.text:
                    for r in p.runs:
                        if k in r.text:
                            r.text = r.text.replace(k, v)
        word_buf = io.BytesIO()
        doc.save(word_buf)
        word_buf.seek(0)

    st.download_button(
        label="📥 SCARICA LETTERA DI TRASMISSIONE (.DOCX)",
        data=word_buf,
        file_name=f"trasmissione_{oggi.strftime('%Y%m%d')}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
        use_container_width=True
    )
